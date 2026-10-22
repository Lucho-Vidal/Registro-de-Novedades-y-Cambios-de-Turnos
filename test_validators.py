"""Unit tests for validators.py - Validación de datos para formularios."""

import pytest
from datetime import datetime
import tkinter as tk
from ttkbootstrap import DateEntry
from validators import (
    parsear_fecha,
    validar_campos_requeridos_novedades,
    validar_campos_requeridos_cambios,
    validar_datos_edicion_novedad,
    validar_datos_edicion_cambio,
)


@pytest.fixture(scope="session")
def tk_root():
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()


class TestParsearFecha:
    """Pruebas para la función parsear_fecha."""

    def test_parsear_fecha_valida(self):
        """Debe parsear una fecha válida en formato DD/MM/YYYY."""
        resultado = parsear_fecha("31/12/2025", "Fecha", obligatorio=False)
        assert resultado == datetime(2025, 12, 31)

    def test_parsear_fecha_valida_obligatorio(self):
        """Debe parsear una fecha válida cuando es obligatorio."""
        resultado = parsear_fecha("01/01/2026", "Fecha", obligatorio=True)
        assert resultado == datetime(2026, 1, 1)

    def test_parsear_fecha_vacia_no_obligatorio(self):
        """Debe retornar None si la fecha está vacía y no es obligatorio."""
        resultado = parsear_fecha("", "Fecha", obligatorio=False)
        assert resultado is None

    def test_parsear_fecha_vacia_con_espacios_no_obligatorio(self):
        """Debe retornar None si la fecha es solo espacios y no es obligatorio."""
        resultado = parsear_fecha("   ", "Fecha", obligatorio=False)
        assert resultado is None

    def test_parsear_fecha_vacia_obligatorio(self):
        """Debe lanzar ValueError si la fecha está vacía y es obligatorio."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha("", "Fecha inicio", obligatorio=True)
        assert "El campo 'Fecha inicio' es obligatorio" in str(excinfo.value)

    def test_parsear_fecha_formato_invalido(self):
        """Debe lanzar ValueError si el formato no es DD/MM/YYYY."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha("2025-12-31", "Fecha", obligatorio=True)
        assert "formato DD/MM/AAAA" in str(excinfo.value)

    def test_parsear_fecha_formato_invalido_mm_dd(self):
        """Debe lanzar ValueError si el formato es MM/DD/YYYY."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha("12/31/2025", "Fecha", obligatorio=True)
        assert "formato DD/MM/AAAA" in str(excinfo.value)

    def test_parsear_fecha_con_espacios_trimmed(self):
        """Debe trimmar espacios antes de parsear."""
        resultado = parsear_fecha("  15/06/2025  ", "Fecha", obligatorio=True)
        assert resultado == datetime(2025, 6, 15)

    def test_parsear_fecha_dia_invalido(self):
        """Debe lanzar ValueError si el día es inválido (ej. 32/01/2025)."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha("32/01/2025", "Fecha", obligatorio=True)
        assert "formato DD/MM/AAAA" in str(excinfo.value)

    def test_parsear_fecha_mes_invalido(self):
        """Debe lanzar ValueError si el mes es inválido (ej. 15/13/2025)."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha("15/13/2025", "Fecha", obligatorio=True)
        assert "formato DD/MM/AAAA" in str(excinfo.value)

    def test_parsear_fecha_valor_none(self):
        """Debe manejar None como valor vacío."""
        resultado = parsear_fecha(None, "Fecha", obligatorio=False)
        assert resultado is None

    def test_parsear_fecha_valor_entero(self):
        """Debe lanzar ValueError si recibe un número."""
        with pytest.raises(ValueError) as excinfo:
            parsear_fecha(15062025, "Fecha", obligatorio=True)
        assert "formato DD/MM/AAAA" in str(excinfo.value)


class MockDateEntry:
    def __init__(self, value="", master=None):
        self.entry = tk.Entry(master=master)
        self.entry.insert(0, value)

    def get(self):
        return self.entry.get()


class TestValidarCamposRequeridosNovedades:
    """Pruebas para validar_campos_requeridos_novedades."""

    @pytest.fixture
    def campos_validos(self, tk_root):
        """Fixture con campos válidos para una novedad."""
        legajo = tk.StringVar(master=tk_root, value="12345")
        apellidos_nombres = tk.StringVar(master=tk_root, value="García, Juan")
        novedad = tk.StringVar(master=tk_root, value="Licencia")
        referencia_estacion = tk.StringVar(master=tk_root, value="EST001")
        supervisor = tk.StringVar(master=tk_root, value="López, Carlos")
        fecha_inicio = MockDateEntry("15/06/2025", master=tk_root)
        fecha_fin = MockDateEntry("20/06/2025", master=tk_root)
        tipo_novedades = ["Licencia", "Enfermedad", "Falta"]

        return {
            "legajo": legajo,
            "apellidos_nombres": apellidos_nombres,
            "novedad": novedad,
            "referencia_estacion": referencia_estacion,
            "supervisor": supervisor,
            "fecha_inicio": fecha_inicio,
            "fecha_fin": fecha_fin,
            "tipo_novedades": tipo_novedades,
        }

    def test_validacion_exitosa(self, campos_validos):
        """Debe retornar (True, None) si todos los campos son válidos."""
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is True
        assert msg is None

    def test_legajo_vacio(self, campos_validos):
        """Debe retornar error si Legajo está vacío."""
        campos_validos["legajo"].set("")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Legajo" in msg

    def test_apellidos_nombres_vacio(self, campos_validos):
        """Debe retornar error si Apellido y Nombre está vacío."""
        campos_validos["apellidos_nombres"].set("")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Apellido y Nombre" in msg

    def test_novedad_vacia(self, campos_validos):
        """Debe retornar error si Tipo de Novedad está vacío."""
        campos_validos["novedad"].set("")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Tipo de Novedad" in msg

    def test_novedad_invalida(self, campos_validos):
        """Debe retornar error si Tipo de Novedad no está en la lista válida."""
        campos_validos["novedad"].set("Tipo Inválido")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "valido" in msg or "Tipo de Novedad" in msg

    def test_referencia_estacion_vacia(self, campos_validos):
        """Debe retornar error si Referencia Estación está vacía."""
        campos_validos["referencia_estacion"].set("")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Referencia Estacion" in msg

    def test_supervisor_vacio(self, campos_validos):
        """Debe retornar error si Supervisor está vacío."""
        campos_validos["supervisor"].set("")
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Supervisor" in msg

    def test_fecha_inicio_invalida(self, campos_validos):
        """Debe retornar error si Fecha de Inicio tiene formato inválido."""
        campos_validos["fecha_inicio"].entry.delete(0, tk.END)
        campos_validos["fecha_inicio"].entry.insert(0, "2025-06-15")  # Formato incorrecto
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "Fecha de inicio novedad" in msg or "formato" in msg

    def test_fecha_fin_anterior_a_inicio(self, campos_validos):
        """Debe retornar error si Fecha Fin es anterior a Fecha Inicio."""
        campos_validos["fecha_fin"].entry.delete(0, tk.END)
        campos_validos["fecha_fin"].entry.insert(0, "10/06/2025")  # Antes que inicio (15/06)
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is False
        assert "no puede ser anterior" in msg

    def test_fecha_fin_vacia_es_valida(self, campos_validos):
        """Debe permitir que Fecha Fin esté vacía (es opcional)."""
        campos_validos["fecha_fin"].entry.delete(0, tk.END)
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is True
        assert msg is None

    def test_fecha_fin_igual_a_inicio(self, campos_validos):
        """Debe permitir que Fecha Fin sea igual a Fecha Inicio."""
        campos_validos["fecha_fin"].entry.delete(0, tk.END)
        campos_validos["fecha_fin"].entry.insert(0, "15/06/2025")  # Igual que inicio
        es_valido, msg = validar_campos_requeridos_novedades(
            campos_validos["legajo"],
            campos_validos["apellidos_nombres"],
            campos_validos["novedad"],
            campos_validos["referencia_estacion"],
            campos_validos["supervisor"],
            campos_validos["fecha_inicio"],
            campos_validos["fecha_fin"],
            campos_validos["tipo_novedades"],
        )
        assert es_valido is True
        assert msg is None


class TestValidarCamposRequeridosCambios:
    """Pruebas para validar_campos_requeridos_cambios."""

    @pytest.fixture
    def campos_validos_cambios(self, tk_root):
        """Fixture con campos válidos para un cambio de turno."""
        legajo = tk.StringVar(master=tk_root, value="12345")
        apellidos_nombres = tk.StringVar(master=tk_root, value="García, Juan")
        legajo_2 = tk.StringVar(master=tk_root, value="67890")
        apellidos_nombres_2 = tk.StringVar(master=tk_root, value="López, Maria")
        fecha_cambio = MockDateEntry("15/06/2025", master=tk_root)
        referencia_estacion = tk.StringVar(master=tk_root, value="EST001")
        supervisor = tk.StringVar(master=tk_root, value="Pérez, Carlos")

        return {
            "legajo": legajo,
            "apellidos_nombres": apellidos_nombres,
            "legajo_2": legajo_2,
            "apellidos_nombres_2": apellidos_nombres_2,
            "fecha_cambio": fecha_cambio,
            "referencia_estacion": referencia_estacion,
            "supervisor": supervisor,
        }

    def test_validacion_exitosa(self, campos_validos_cambios):
        """Debe retornar (True, None) si todos los campos son válidos."""
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is True
        assert msg is None

    def test_legajo_vacio(self, campos_validos_cambios):
        """Debe retornar error si Legajo 1 está vacío."""
        campos_validos_cambios["legajo"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Legajo" in msg

    def test_apellidos_nombres_vacio(self, campos_validos_cambios):
        """Debe retornar error si Apellido y Nombre 1 está vacío."""
        campos_validos_cambios["apellidos_nombres"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Apellido y Nombre" in msg

    def test_legajo_2_vacio(self, campos_validos_cambios):
        """Debe retornar error si Legajo 2 está vacío."""
        campos_validos_cambios["legajo_2"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Legajo 2" in msg

    def test_apellidos_nombres_2_vacio(self, campos_validos_cambios):
        """Debe retornar error si Apellido y Nombre 2 está vacío."""
        campos_validos_cambios["apellidos_nombres_2"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Apellido y Nombre 2" in msg

    def test_fecha_cambio_vacia(self, campos_validos_cambios):
        """Debe retornar error si Fecha de cambio de turno está vacía."""
        campos_validos_cambios["fecha_cambio"].entry.delete(0, tk.END)
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Fecha de cambio de turno" in msg

    def test_fecha_cambio_formato_invalido(self, campos_validos_cambios):
        """Debe retornar error si Fecha de cambio tiene formato inválido."""
        campos_validos_cambios["fecha_cambio"].entry.delete(0, tk.END)
        campos_validos_cambios["fecha_cambio"].entry.insert(0, "2025-06-15")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False

    def test_referencia_estacion_vacia(self, campos_validos_cambios):
        """Debe retornar error si Referencia Estación está vacía."""
        campos_validos_cambios["referencia_estacion"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Referencia Estacion" in msg

    def test_supervisor_vacio(self, campos_validos_cambios):
        """Debe retornar error si Supervisor está vacío."""
        campos_validos_cambios["supervisor"].set("")
        es_valido, msg = validar_campos_requeridos_cambios(
            campos_validos_cambios["legajo"],
            campos_validos_cambios["apellidos_nombres"],
            campos_validos_cambios["legajo_2"],
            campos_validos_cambios["apellidos_nombres_2"],
            campos_validos_cambios["fecha_cambio"],
            campos_validos_cambios["referencia_estacion"],
            campos_validos_cambios["supervisor"],
        )
        assert es_valido is False
        assert "Supervisor" in msg


class TestValidarDatosEdicion:
    """Pruebas para validación de datos en el editor de registros."""

    def test_edicion_novedad_valida(self):
        data = {
            "legajo": "1234",
            "apellidos_nombres": "Pérez, Juan",
            "novedad": "Licencia",
            "referencia_estacion": "Retiro (PC)",
            "supervisor": "Gómez",
            "fecha_inicio": "10/01/2026",
            "fecha_fin": "15/01/2026",
        }
        valido, err = validar_datos_edicion_novedad(
            data,
            tipo_novedades=["Licencia", "Enfermedad"],
            personal_estacion=["Retiro (PC)", "Constitución (TY)"],
        )
        assert valido is True
        assert err is None

    def test_edicion_novedad_fecha_fin_invalida(self):
        data = {
            "legajo": "1234",
            "apellidos_nombres": "Pérez, Juan",
            "novedad": "Licencia",
            "referencia_estacion": "Retiro (PC)",
            "supervisor": "Gómez",
            "fecha_inicio": "10/01/2026",
            "fecha_fin": "05/01/2026",  # Anterior
        }
        valido, err = validar_datos_edicion_novedad(data)
        assert valido is False
        assert "no puede ser anterior" in err

    def test_edicion_novedad_tipo_invalido(self):
        data = {
            "legajo": "1234",
            "apellidos_nombres": "Pérez, Juan",
            "novedad": "Tipo Inexistente",
            "referencia_estacion": "Retiro (PC)",
            "supervisor": "Gómez",
            "fecha_inicio": "10/01/2026",
        }
        valido, err = validar_datos_edicion_novedad(data, tipo_novedades=["Licencia"])
        assert valido is False
        assert "Tipo de Novedad" in err

    def test_edicion_cambio_valido(self):
        data = {
            "legajo_1": "111",
            "apellidos_nombres_1": "Pérez, Juan",
            "legajo_2": "222",
            "apellidos_nombres_2": "López, Ana",
            "referencia_estacion": "Retiro (PC)",
            "supervisor": "Gómez",
            "fecha_cambio": "20/02/2026",
        }
        valido, err = validar_datos_edicion_cambio(data, personal_estacion=["Retiro (PC)"])
        assert valido is True
        assert err is None

    def test_edicion_cambio_legajo_invalido(self):
        data = {
            "legajo_1": "abc",
            "apellidos_nombres_1": "Pérez, Juan",
            "legajo_2": "222",
            "apellidos_nombres_2": "López, Ana",
            "referencia_estacion": "Retiro (PC)",
            "supervisor": "Gómez",
            "fecha_cambio": "20/02/2026",
        }
        valido, err = validar_datos_edicion_cambio(data)
        assert valido is False
        assert "entero" in err


class TestNegocioYAutenticacion:
    """Pruebas de negocio: atomicidad de reseteo de clave, edición de fecha_fin y métricas de dashboard."""

    def test_resetear_password_atomico_falla_envio(self, tmp_path):
        from sqlite_store import SQLiteStore
        from auth import AuthService

        db_file = tmp_path / "test.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        auth = AuthService(store)

        auth.crear_administrador_inicial("admin", "clave_original", 100)
        user_antes = auth.autenticar("admin", "clave_original")
        assert user_antes is not None
        hash_antes = user_antes["password_hash"]

        def mock_fallo_envio(email, temporal):
            raise RuntimeError("Error al conectar con Outlook")

        with pytest.raises(RuntimeError):
            auth.resetear_password("admin", send_fn=mock_fallo_envio)

        # La contraseña en SQLite NO debe haberse modificado
        user_despues = auth.autenticar("admin", "clave_original")
        assert user_despues is not None
        assert user_despues["password_hash"] == hash_antes

    def test_actualizar_novedad_sin_fecha_fin_permite_edicion_fuera_de_ventana(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService

        db_file = tmp_path / "test_records.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        # Insertar novedad con fecha de registro de hace 10 días (ventana de 24hs superada) y sin fecha fin
        fecha_antigua = "2026-01-01 10:00:00"
        with store.write_transaction() as conn:
            cursor = conn.execute(
                """INSERT INTO novedades
                   (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion,
                    turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor)
                   VALUES (1, ?, 100, 'Test, Juan', 'Cond', 'PC', 'T1', 'Lun', 'Licencia', '01/01/2026', NULL, 'Retiro', 'Sup')""",
                (fecha_antigua,)
            )

        # Intentar actualizar_novedad completa fuera de ventana debe fallar aunque falte fecha_fin
        data = {
            "legajo": 100,
            "apellidos_nombres": "Test, Juan",
            "especialidad": "Cond",
            "dotacion": "PC",
            "turnos": "T1",
            "franco": "Lun",
            "novedad": "Licencia",
            "fecha_inicio": "01/01/2026",
            "fecha_fin": "15/01/2026",
            "referencia_estacion": "Retiro",
            "supervisor": "Sup",
            "observaciones": "Fin completado",
        }
        with pytest.raises(ValueError, match="El registro superó el tiempo permitido para editar"):
            records.actualizar_novedad(1, data, bypass_window=False)

        # actualizar_fecha_fin_novedad sí debe funcionar fuera de ventana y concatenar observaciones
        records.actualizar_fecha_fin_novedad(1, "15/01/2026", nueva_observacion="Alta médica presentada")

        fila = records.obtener_novedad(1)
        assert fila["fecha_fin"] == "15/01/2026"
        assert "Alta médica presentada" in fila["observaciones"]

        # Intentar actualizar fecha_fin nuevamente debe fallar porque ya tiene fecha fin
        with pytest.raises(ValueError, match="El registro ya cuenta con una fecha de fin"):
            records.actualizar_fecha_fin_novedad(1, "20/01/2026")

    def test_actualizar_fecha_fin_casos_borde(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService

        db_file = tmp_path / "test_fecha_fin_edge.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        with store.write_transaction() as conn:
            conn.execute(
                """INSERT INTO novedades
                   (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion,
                    turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, observaciones)
                   VALUES (1, '2026-01-01 10:00:00', 101, 'Gomez, Ana', 'Guarda', 'TY', 'T2', 'Mar', 'Enfermedad', '10/01/2026', NULL, 'Constitución', 'Sup', 'Nota inicial')"""
            )

        # 1. Fecha fin vacía debe fallar
        with pytest.raises(ValueError, match="Debe ingresar la fecha de fin"):
            records.actualizar_fecha_fin_novedad(1, "")

        # 2. Formato inválido debe fallar
        with pytest.raises(ValueError, match="debe tener formato DD/MM/AAAA"):
            records.actualizar_fecha_fin_novedad(1, "2026-01-20")

        # 3. Fecha fin anterior a fecha inicio (10/01/2026) debe fallar
        with pytest.raises(ValueError, match="no puede ser anterior a la fecha de inicio"):
            records.actualizar_fecha_fin_novedad(1, "05/01/2026")

        # 4. Concatenación con texto previo existente
        records.actualizar_fecha_fin_novedad(1, "12/01/2026", nueva_observacion="Segunda nota")
        fila = records.obtener_novedad(1)
        assert fila["fecha_fin"] == "12/01/2026"
        assert fila["observaciones"] == "Nota inicial\nSegunda nota"

    def test_obtener_metricas_dashboard_unificado(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService
        from datetime import datetime

        db_file = tmp_path / "test_dash.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        ahora_iso = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with store.write_transaction() as conn:
            conn.execute(
                """INSERT INTO novedades
                   (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion,
                    turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor)
                   VALUES (1, ?, 100, 'Test', 'Cond', 'PC', 'T1', 'Lun', 'Licencia', '01/01/2026', NULL, 'Retiro', 'Sup')""",
                (ahora_iso,)
            )

        metricas = records.obtener_metricas_dashboard(30)
        assert metricas["resumen"]["novedades_hoy"] >= 1
        assert len(metricas["por_tipo"]) >= 1
        assert metricas["por_tipo"][0][0] == "Licencia"
        assert len(metricas["tendencia"]) == 30


class TestTtkbootstrapPatch:
    """Pruebas de robustez para el parche de ttkbootstrap contra 'bad window path name'."""

    def test_cambio_tema_con_combobox_destruido_no_falla(self, tk_root):
        from ttk_patch import aplicar_patch_ttkbootstrap
        import ttkbootstrap as tb
        from tkinter import ttk

        aplicar_patch_ttkbootstrap()
        style = tb.Style("flatly")

        # Crear ventana secundaria con combobox y luego destruirla
        top = tk.Toplevel(tk_root)
        cb = ttk.Combobox(top, values=["A", "B", "C"])
        cb.pack()
        top.update_idletasks()
        top.destroy()
        tk_root.update_idletasks()

        # Cambiar temas sucesivamente no debe levantar TclError ni crashear
        style.theme_use("darkly")
        style.theme_use("flatly")


class TestTurnosSinCubrirYTiposNovedad:
    """Pruebas para tipos de novedad (dias, exportar_descubierto) y exportación de turnos sin cubrir."""

    def test_tipos_novedad_crud_y_dict(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService

        db_file = tmp_path / "test_tipos.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        diccionario_inicial = store.get_tipo_novedades_dict()
        assert "Dona sangre" in diccionario_inicial
        assert diccionario_inicial["Dona sangre"]["dias"] == 1
        assert diccionario_inicial["Dona sangre"]["exportar_descubierto"] == 1

        assert "Informe" in diccionario_inicial
        assert diccionario_inicial["Informe"]["dias"] == 1
        assert diccionario_inicial["Informe"]["exportar_descubierto"] == 0

        # Crear nuevo tipo
        tipo_id = records.crear_tipo("Licencia Examen", dias=3, exportar_descubierto=1)

        diccionario = store.get_tipo_novedades_dict()
        assert "Licencia Examen" in diccionario
        assert diccionario["Licencia Examen"]["dias"] == 3
        assert diccionario["Licencia Examen"]["exportar_descubierto"] == 1

        # Actualizar tipo
        records.actualizar_tipo(tipo_id, "Licencia Examen Final", dias=4, exportar_descubierto=0, activo=1)
        tipos = records.listar_tipos(incluir_inactivos=True)
        actualizado = [t for t in tipos if t[0] == tipo_id][0]
        assert actualizado[1] == "Licencia Examen Final"
        assert actualizado[2] == 4
        assert actualizado[3] == 0

    def test_obtener_turnos_sin_cubrir_filtrado(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService
        from datetime import datetime

        db_file = tmp_path / "test_sin_cubrir.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        # Configurar tipos (Informe y Dona sangre ya vienen sembrados con exportar_descubierto=0 y 1 respectivamente)
        records.crear_tipo("Enfermedad", dias=0, exportar_descubierto=1)

        ahora_iso = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with store.write_transaction() as conn:
            # 1. Novedad que cubre la fecha (05/03/2026 a 10/03/2026) - DEBE INCLUIRSE para 08/03/2026
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (1, ?, 101, 'Perez, Juan', 'Guarda', 'PC', 'T1', 'Lun', 'Enfermedad', '05/03/2026', '10/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 2. Novedad sin fecha fin (inicio 01/03/2026) - DEBE INCLUIRSE para 08/03/2026
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (2, ?, 102, 'Gomez, Ana', 'Conductor', 'TY', 'T2', 'Mar', 'Enfermedad', '01/03/2026', NULL, 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 3. Novedad de 1 solo día exacta para 08/03/2026 - DEBE INCLUIRSE
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (3, ?, 103, 'Lopez, Carlos', 'Guarda', 'PC', 'T1', 'Mie', 'Dona sangre', '08/03/2026', '08/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 4. Novedad con exportar_descubierto=0 (Informe) - NO DEBE INCLUIRSE
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (4, ?, 104, 'Diaz, Maria', 'Guarda', 'PC', 'T1', 'Jue', 'Informe', '08/03/2026', '08/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 5. Novedad fuera de rango (terminó el 07/03/2026) - NO DEBE INCLUIRSE para 08/03/2026
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (5, ?, 105, 'Sosa, Luis', 'Guarda', 'PC', 'T1', 'Vie', 'Enfermedad', '01/03/2026', '07/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 6. Novedad futura (inicia 09/03/2026) - NO DEBE INCLUIRSE para 08/03/2026
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (6, ?, 106, 'Ruiz, Pedro', 'Guarda', 'PC', 'T1', 'Sab', 'Enfermedad', '09/03/2026', '15/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )
            # 7. Novedad eliminada (activo=0) - NO DEBE INCLUIRSE
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (7, ?, 107, 'Vazquez, Raul', 'Guarda', 'PC', 'T1', 'Dom', 'Enfermedad', '05/03/2026', '10/03/2026', 'Retiro', 'Sup', 0)""",
                (ahora_iso,)
            )

        # Consulta general para 08/03/2026 (sin filtro de dotación)
        filas = records.obtener_turnos_sin_cubrir("08/03/2026")
        ids = [f["id"] for f in filas]
        assert set(ids) == {1, 2, 3}

        # Consulta filtrando por dotación "TY"
        filas_ty = records.obtener_turnos_sin_cubrir("08/03/2026", dotacion="TY")
        assert len(filas_ty) == 1
        assert filas_ty[0]["id"] == 2
        assert filas_ty[0]["dotacion"] == "TY"

        # Consulta filtrando por dotación "PC"
        filas_pc = records.obtener_turnos_sin_cubrir("08/03/2026", dotacion="PC")
        assert len(filas_pc) == 2
        assert set(f["id"] for f in filas_pc) == {1, 3}

    def test_exportar_turnos_sin_cubrir_excel(self, tmp_path):
        import openpyxl
        from sqlite_store import SQLiteStore
        from records_service import RecordsService
        from excel_exporter import export_turnos_sin_cubrir
        from datetime import datetime

        db_file = tmp_path / "test_export_desc.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)

        records.crear_tipo("Enfermedad", dias=0, exportar_descubierto=1)
        ahora_iso = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with store.write_transaction() as conn:
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (1, ?, 999, 'Perez, Juan', 'Guarda', 'PC', 'T1', 'Lun', 'Enfermedad', '10/02/2026', '15/02/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,)
            )

        output_xlsx = tmp_path / "turnos_sin_cubrir.xlsx"
        export_turnos_sin_cubrir(records, str(output_xlsx), "12/02/2026")

        assert output_xlsx.exists()
        wb = openpyxl.load_workbook(str(output_xlsx))
        ws = wb.active
        assert ws.title == "Turnos Sin Cubrir"
        headers = [cell.value for cell in ws[1]]
        assert "LEGAJO" in headers
        assert "APELLIDOS Y NOMBRES" in headers
        assert "NOVEDAD" in headers
        row2 = [cell.value for cell in ws[2]]
        assert 999 in row2 or "999" in row2
        assert "Perez, Juan" in row2

    def test_calcular_y_asignar_fecha_fin_forms(self, tk_root):
        from forms import FormsManager
        from unittest.mock import MagicMock

        app = MagicMock()
        app.novedad_var = tk.StringVar(tk_root)
        app.fecha_inicio_novedad_var = tk.StringVar(tk_root)
        app.fecha_fin_novedad_var = tk.StringVar(tk_root)
        app.tipo_novedades_dict = {
            "Dona sangre": {"dias": 1, "exportar_descubierto": 1},
            "Licencia 3 dias": {"dias": 3, "exportar_descubierto": 1},
            "Enfermedad": {"dias": 0, "exportar_descubierto": 1},
        }

        entry_inicio = tk.Entry(tk_root)
        entry_fin = tk.Entry(tk_root)

        app.fecha_inicio_novedad_entry = MagicMock()
        app.fecha_inicio_novedad_entry.entry = entry_inicio

        app.fecha_fin_novedad_entry = MagicMock()
        app.fecha_fin_novedad_entry.entry = entry_fin

        forms = FormsManager(app)
        forms.edicion = None

        # 1. Caso días = 1: Dona sangre -> fecha fin igual a fecha inicio
        entry_inicio.delete(0, tk.END)
        entry_inicio.insert(0, "15/03/2026")
        entry_fin.delete(0, tk.END)
        app.novedad_var.set("Dona sangre")

        forms._calcular_y_asignar_fecha_fin(origen="tipo")
        assert entry_fin.get() == "15/03/2026"
        assert app.fecha_fin_novedad_var.get() == "15/03/2026"

        # 2. Caso días = 3: Licencia 3 dias -> 15/03/2026 + 2 = 17/03/2026
        entry_fin.delete(0, tk.END)
        app.novedad_var.set("Licencia 3 dias")

        forms._calcular_y_asignar_fecha_fin(origen="tipo")
        assert entry_fin.get() == "17/03/2026"
        assert app.fecha_fin_novedad_var.get() == "17/03/2026"

        # 3. Caso días = 0: cambio a Enfermedad -> LIMPIA el input de fecha fin si viene de cambio de tipo
        app.novedad_var.set("Enfermedad")
        forms._calcular_y_asignar_fecha_fin(origen="tipo")
        assert entry_fin.get() == ""
        assert app.fecha_fin_novedad_var.get() == ""

        # 4. Caso días = 0 con fecha_inicio modificada: no debe borrar una fecha ingresada manualmente
        entry_fin.insert(0, "20/03/2026")
        app.fecha_fin_novedad_var.set("20/03/2026")
        entry_inicio.delete(0, tk.END)
        entry_inicio.insert(0, "16/03/2026")
        forms._calcular_y_asignar_fecha_fin(origen="fecha_inicio")
        assert entry_fin.get() == "20/03/2026"
        assert app.fecha_fin_novedad_var.get() == "20/03/2026"


class TestReemplazosYPersonalCiclo:
    """Pruebas para asignación de reemplazos CICLO, no solapamiento, estado de cobertura,
    vista de Personal de Ciclo y calendario mensual.
    """

    def test_es_turno_ciclo_y_validar_tramos_reemplazo(self):
        import validators
        assert validators.es_turno_ciclo("CICLO 1") is True
        assert validators.es_turno_ciclo("ciclo mañana") is True
        assert validators.es_turno_ciclo("  Ciclo  ") is True
        assert validators.es_turno_ciclo("06:00 A 14:00") is False
        assert validators.es_turno_ciclo("") is False

        # Titular con turno CICLO no puede tener reemplazos
        ok, err = validators.validar_tramos_reemplazo(
            "01/03/2026", "10/03/2026", "CICLO 2",
            [{"legajo_reemplazo": 500, "apellidos_nombres_reemplazo": "Relevo Uno", "fecha_inicio": "01/03/2026", "fecha_fin": "10/03/2026"}],
        )
        assert ok is False
        assert "CICLO" in err

        # Dos tramos secuenciales sin solaparse dentro del rango -> Válido
        tramos_ok = [
            {"legajo_reemplazo": 501, "apellidos_nombres_reemplazo": "Relevo Uno", "fecha_inicio": "01/03/2026", "fecha_fin": "05/03/2026"},
            {"legajo_reemplazo": 502, "apellidos_nombres_reemplazo": "Relevo Dos", "fecha_inicio": "06/03/2026", "fecha_fin": "10/03/2026"},
        ]
        ok, err = validators.validar_tramos_reemplazo("01/03/2026", "10/03/2026", "T1", tramos_ok)
        assert ok is True
        assert err is None

        # Dos tramos solapados -> Inválido
        tramos_solapados = [
            {"legajo_reemplazo": 501, "apellidos_nombres_reemplazo": "Relevo Uno", "fecha_inicio": "01/03/2026", "fecha_fin": "06/03/2026"},
            {"legajo_reemplazo": 502, "apellidos_nombres_reemplazo": "Relevo Dos", "fecha_inicio": "06/03/2026", "fecha_fin": "10/03/2026"},
        ]
        ok, err = validators.validar_tramos_reemplazo("01/03/2026", "10/03/2026", "T1", tramos_solapados)
        assert ok is False
        assert "solapa" in err.lower()

        # Tramo fuera del rango de la novedad -> Inválido
        tramos_fuera = [
            {"legajo_reemplazo": 501, "apellidos_nombres_reemplazo": "Relevo Uno", "fecha_inicio": "01/03/2026", "fecha_fin": "12/03/2026"},
        ]
        ok, err = validators.validar_tramos_reemplazo("01/03/2026", "10/03/2026", "T1", tramos_fuera)
        assert ok is False
        assert "superar" in err.lower()

    def test_flujo_completo_reemplazos_cobertura_y_calendario(self, tmp_path):
        import openpyxl
        from sqlite_store import SQLiteStore
        from records_service import RecordsService
        from excel_exporter import export_database, export_turnos_sin_cubrir
        from datetime import datetime

        db_file = tmp_path / "test_ciclo.sqlite"
        store = SQLiteStore(str(db_file))
        store.initialize()
        records = RecordsService(store)
        records.crear_tipo("Enfermedad", dias=0, exportar_descubierto=1)
        ahora_iso = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Crear empleados: titulares y personal de CICLO
        for leg, nom, esp, dot, tur, fra in (
            (100, "Titular Uno", "Guarda", "PC", "06:00 A 14:00", "Domingo"),
            (101, "Titular Dos", "Guarda", "PC", "14:00 A 22:00", "Sabado"),
            (201, "Ciclo Gomez", "Guarda", "PC", "CICLO 1", "Lunes"),
            (202, "Ciclo Perez", "Guarda", "PC", "CICLO 2", "Martes"),
            (203, "Ciclo OtraEsp", "Conductor", "PC", "CICLO 3", "Miercoles"),
            (204, "Ciclo OtraDot", "Guarda", "TY", "CICLO 4", "Jueves"),
        ):
            records.crear_empleado({
                "legajo": leg,
                "apellidos_nombres": nom,
                "especialidad": esp,
                "dotacion": dot,
                "turnos": tur,
                "franco": fra,
            })

        # 1. Filtrado de candidatos CICLO: misma dotación (PC), misma especialidad (Guarda) y turno CICLO*
        candidatos = records.obtener_candidatos_ciclo("PC", "Guarda")
        legajos_cand = [c["legajo"] for c in candidatos]
        assert set(legajos_cand) == {201, 202}

        # 2. Insertar dos novedades para titulares
        with store.write_transaction() as conn:
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (1, ?, 100, 'Titular Uno', 'Guarda', 'PC', '06:00 A 14:00', 'Domingo', 'Enfermedad', '01/03/2026', '10/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,),
            )
            conn.execute(
                """INSERT INTO novedades (id, registrado_en, legajo, apellidos_nombres, especialidad, dotacion, turnos, franco, novedad, fecha_inicio, fecha_fin, referencia_estacion, supervisor, activo)
                   VALUES (2, ?, 101, 'Titular Dos', 'Guarda', 'PC', '14:00 A 22:00', 'Sabado', 'Enfermedad', '04/03/2026', '08/03/2026', 'Retiro', 'Sup', 1)""",
                (ahora_iso,),
            )

        # Inicialmente novedad #1 no está cubierta
        est, det = records.calcular_estado_cobertura(1, "01/03/2026", "10/03/2026", "06:00 A 14:00")
        assert est == "No"

        # 3. Asignar cobertura parcial (01/03 al 05/03 con 201)
        records.guardar_reemplazos_novedad(
            1,
            [{"legajo_reemplazo": 201, "apellidos_nombres_reemplazo": "Ciclo Gomez", "turnos_reemplazo": "CICLO 1", "franco_reemplazo": "Lunes", "fecha_inicio": "01/03/2026", "fecha_fin": "05/03/2026"}],
        )
        est, det = records.calcular_estado_cobertura(1, "01/03/2026", "10/03/2026", "06:00 A 14:00")
        assert est == "Parcial"
        assert "Ciclo Gomez" in det

        # 4. Intentar asignar a Ciclo Gomez (201) en la novedad #2 (04/03 al 08/03) -> debe fallar por solapamiento
        with pytest.raises(ValueError, match="ya reemplaza a"):
            records.guardar_reemplazos_novedad(
                2,
                [{"legajo_reemplazo": 201, "apellidos_nombres_reemplazo": "Ciclo Gomez", "turnos_reemplazo": "CICLO 1", "franco_reemplazo": "Lunes", "fecha_inicio": "04/03/2026", "fecha_fin": "08/03/2026"}],
            )

        # 5. Completar la cobertura de la novedad #1 con Ciclo Perez (202) del 06/03 al 10/03 -> Cobertura total ("Sí")
        records.guardar_reemplazos_novedad(
            1,
            [
                {"legajo_reemplazo": 201, "apellidos_nombres_reemplazo": "Ciclo Gomez", "turnos_reemplazo": "CICLO 1", "franco_reemplazo": "Lunes", "fecha_inicio": "01/03/2026", "fecha_fin": "05/03/2026"},
                {"legajo_reemplazo": 202, "apellidos_nombres_reemplazo": "Ciclo Perez", "turnos_reemplazo": "CICLO 2", "franco_reemplazo": "Martes", "fecha_inicio": "06/03/2026", "fecha_fin": "10/03/2026"},
            ],
        )
        est, det = records.calcular_estado_cobertura(1, "01/03/2026", "10/03/2026", "06:00 A 14:00")
        assert est == "Sí"

        # 6. Estado de Personal de Ciclo el 03/03/2026 (201 ocupado reemplazando a 100, 202 libre)
        estado_ciclo = records.obtener_estado_personal_ciclo("03/03/2026", dotacion="PC")
        legajos_ocupados = [e["legajo"] for e in estado_ciclo["ocupados"]]
        legajos_libres = [e["legajo"] for e in estado_ciclo["libres"]]
        assert 201 in legajos_ocupados
        assert 202 in legajos_libres
        assert estado_ciclo["ocupados"][0]["reemplazando_legajo"] == 100

        # 7. Calendario mensual de Ciclo Gomez (201) para Marzo 2026 (31 días)
        cal = records.obtener_calendario_mensual_ciclo(201, 2026, 3)
        assert len(cal) == 31
        # Del 01/03 al 05/03 reemplaza a Titular Uno (legajo 100)
        for d in cal[:5]:
            assert d["reemplazando_legajo"] == 100
            assert "Titular Uno" in d["reemplazando_nombre"]
        # El 06/03 ya no tiene reemplazo asignado
        assert cal[5]["reemplazando_legajo"] is None

        # 8. Exportaciones incluyen CUBIERTO y REEMPLAZO(S)
        out_sin_cubrir = tmp_path / "sin_cubrir_cob.xlsx"
        export_turnos_sin_cubrir(records, str(out_sin_cubrir), "03/03/2026", dotacion="PC")
        wb_sc = openpyxl.load_workbook(str(out_sin_cubrir))
        ws_sc = wb_sc["Turnos Sin Cubrir"]
        headers_sc = [c.value for c in ws_sc[1]]
        assert "CUBIERTO" in headers_sc
        assert "REEMPLAZO" in headers_sc
        vals_sc = [c.value for c in ws_sc[2]]
        assert "Sí" in vals_sc

        out_db = tmp_path / "db_export.xlsx"
        export_database(store, str(out_db), tables=["NOVEDADES"])
        wb_db = openpyxl.load_workbook(str(out_db))
        ws_db = wb_db["NOVEDADES"]
        headers_db = [c.value for c in ws_db[1]]
        assert "CUBIERTO" in headers_db
        assert "REEMPLAZOS" in headers_db

    def test_actualizar_fecha_fin_solo_observacion(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService

        db_path = tmp_path / "test_obs.sqlite"
        store = SQLiteStore(str(db_path))
        store.initialize()
        records = RecordsService(store)

        nov_id = store.insert_novedad([
            1, "2026-03-01 10:00:00", 100, "Titular Uno", "CONDUCTOR", "PC",
            "06:00 A 14:00", "Lunes", "Enfermedad", "01/03/2026", None,
            "Plaza C", "Sup Uno", "Obs inicial", "win_user", None
        ])

        # Actualizar sólo observación sin poner fecha fin
        records.actualizar_fecha_fin_novedad(nov_id, "", nueva_observacion="Seguimiento médico pendiente")
        nov = records.obtener_novedad(nov_id)
        assert nov["fecha_fin"] is None
        assert "Obs inicial" in nov["observaciones"]
        assert "Seguimiento médico pendiente" in nov["observaciones"]

        # Si no se envía ni fecha fin ni nueva observación, debe fallar
        with pytest.raises(ValueError, match="Debe ingresar la fecha de fin o una observación"):
            records.actualizar_fecha_fin_novedad(nov_id, "", nueva_observacion="")

    def test_empleados_franco_horario_y_filtros_ciclo(self, tmp_path):
        from sqlite_store import SQLiteStore
        from records_service import RecordsService

        db_path = tmp_path / "test_ciclo.sqlite"
        store = SQLiteStore(str(db_path))
        store.initialize()
        records = RecordsService(store)

        # 1. Crear empleado ciclo con franco horario
        emp_id = records.crear_empleado({
            "legajo": 501,
            "apellidos_nombres": "Ciclo Conductor",
            "especialidad": "Conductor Eléctrico",
            "dotacion": "PC",
            "turnos": "CICLO 1",
            "franco": "Domingo",
            "franco_horario": "14:00",
        })

        emp_guarda_id = records.crear_empleado({
            "legajo": 502,
            "apellidos_nombres": "Ciclo Guarda",
            "especialidad": "Guarda Eléctrico",
            "dotacion": "PC",
            "turnos": "CICLO 2",
            "franco": "Lunes",
            "franco_horario": "22:00",
        })

        base_rows = store.get_base_rows()
        # Verificar que get_base_rows contiene franco_horario
        row_501 = [r for r in base_rows if r[0] == 501][0]
        assert row_501[6] == "14:00"

        # 2. Actualizar empleado y su franco horario
        records.actualizar_empleado(emp_id, {
            "legajo": 501,
            "apellidos_nombres": "Ciclo Conductor",
            "especialidad": "Conductor Eléctrico",
            "dotacion": "PC",
            "turnos": "CICLO 1",
            "franco": "Domingo",
            "franco_horario": "16:00",
        })
        base_rows2 = store.get_base_rows()
        row_501_act = [r for r in base_rows2 if r[0] == 501][0]
        assert row_501_act[6] == "16:00"

        # 3. Filtrar por especialidad en personal de ciclo
        res_conductor = records.obtener_estado_personal_ciclo("01/03/2026", dotacion="PC", especialidad="Conductor Eléctrico")
        libres_conductores = [e["legajo"] for e in res_conductor["libres"]]
        assert 501 in libres_conductores
        assert 502 not in libres_conductores

        res_guarda = records.obtener_estado_personal_ciclo("01/03/2026", dotacion="PC", especialidad="Guarda Eléctrico")
        libres_guardas = [e["legajo"] for e in res_guarda["libres"]]
        assert 502 in libres_guardas
        assert 501 not in libres_guardas

        # 4. Calendario mensual: verificar que el domingo muestra Franco (desde 16:00) y novedad_id está presente al reemplazar
        # Insertar una novedad y asignarle reemplazo del 02/03 al 05/03 a 501
        nov_id = store.insert_novedad([
            10, "2026-03-01 10:00:00", 999, "Titular Conductor", "Conductor Eléctrico", "PC",
            "06:00 A 14:00", "Lunes", "Enfermedad", "02/03/2026", "05/03/2026",
            "Plaza C", "Sup", "Obs", "win", None
        ])
        records.guardar_reemplazos_novedad(
            10,
            [{"legajo_reemplazo": 501, "apellidos_nombres_reemplazo": "Ciclo Conductor", "turnos_reemplazo": "CICLO 1", "franco_reemplazo": "Domingo", "fecha_inicio": "02/03/2026", "fecha_fin": "05/03/2026"}]
        )

        cal = records.obtener_calendario_mensual_ciclo(501, 2026, 3)
        # 01/03/2026 fue Domingo -> no reemplaza, debe ser Franco (desde 16:00)
        assert cal[0]["fecha"] == "01/03/2026"
        assert cal[0]["estado"] == "Franco (desde 16:00)"
        assert cal[0]["novedad_id"] == ""

        # 02/03/2026 al 05/03/2026 -> Reemplazando con novedad_id 10
        for d in cal[1:5]:
            assert d["estado"] == "Reemplazando"
            assert d["novedad_id"] == 10
            assert d["reemplazando_legajo"] == 999




